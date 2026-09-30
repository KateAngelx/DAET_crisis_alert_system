"""Apply approved Ch 1–3 revisions to match CONNECT-DAET / DAET TOURISM system."""
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = Path(r"C:\Users\kate angel\Downloads\DAET-Tourism-Crisis-Communication-Chap-1-3 2.docx")
BACKUP = SRC.with_name(SRC.stem + ".backup.docx")
OUT = SRC

REPLACEMENTS = [
    ("CHAPTER 3TECHNICAL BACKGROUND", "CHAPTER 3 TECHNICAL BACKGROUND"),
    ("Spring Development", "Sprint Development"),
    ("Spring Review", "Sprint Review"),
    ("Spring Retrospective", "Sprint Retrospective"),
    (
        "compose emergency announcements, categorize incidents, monitor notification history",
        "compose and broadcast emergency announcements, review incident reports, and monitor notification history",
    ),
    (
        "through SMS, email, and web notifications while supporting incident reporting",
        "through SMS, email, and in-app notifications while supporting incident reporting",
    ),
    (
        "multi-channel notifications (SMS, email, and web notifications), emergency alert management",
        "multi-channel notifications (SMS, email, and in-app notifications), emergency alert management",
    ),
    (
        "Although the proposed study does not incorporate location-based services, the effectiveness of digital notifications supports the use of SMS, email, and web notifications for emergency communication.",
        "Although the proposed study does not incorporate continuous real-time GPS tracking of tourists, it provides map-based display of affected areas, route advisories, and hazard locations. The effectiveness of digital notifications supports the use of SMS, email, and in-app notifications for emergency communication.",
    ),
    (
        "To design and develop a web-based Daet Tourism Crisis Communication and Emergency Alert System with features including tourist registration, emergency alert management, user messaging, and multi-channel notification through SMS, email, and web notifications.",
        "To design and develop a web-based Daet Tourism Crisis Communication and Emergency Alert System with features including tourist registration, role-based access for tourists, tourism guides, and administrators, emergency alert management through the Command Center, incident reporting and response tracking, route and hazard advisories, and multi-channel notification through SMS, email, and in-app notifications.",
    ),
    (
        "To implement a centralized database for managing tourist information, emergency alerts, notification records, and user messages to improve emergency communication and information management.",
        "To implement a centralized database for managing tourist and guide profiles, emergency alerts, incident reports, notification and delivery records, route advisories, and area hazards to improve emergency communication and information management.",
    ),
    (
        "The general objective of this study is to develop a Daet Tourism Crisis Communication and Emergency Alert System that enables the Daet Tourism Office to efficiently disseminate timely and reliable emergency notifications to tourists through a centralized web-based platform.",
        "The general objective of this study is to develop a Daet Tourism Crisis Communication and Emergency Alert System that enables the Daet Municipal Tourism Office to efficiently disseminate timely and reliable emergency notifications to tourists and tourism guides through a centralized web-based platform.",
    ),
    (
        "This study covers the analysis, design, development, implementation, and evaluation of the Daet Tourism Crisis Communication and Emergency Alert System, a web-based platform developed for the Daet Tourism Office of the Municipality of Daet, Camarines Norte. The proposed system aims to enhance emergency communication by providing a centralized platform for disseminating timely and reliable emergency notifications to registered tourists.",
        "This study covers the analysis, design, development, implementation, and evaluation of the Daet Tourism Crisis Communication and Emergency Alert System (public brand: DAET TOURISM), a web-based platform developed for the Daet Municipal Tourism Office of the Municipality of Daet, Camarines Norte. The proposed system aims to enhance emergency communication by providing a centralized platform for disseminating timely and reliable emergency notifications to registered tourists and tourism guides, while allowing the public to view active official alerts without an account.",
    ),
    (
        "The system enables authorized personnel to create, manage, and distribute emergency alerts through Short Message Service (SMS), electronic mail (Email), and web notifications. It includes functionalities such as tourist registration, secure user authentication, profile management, emergency alert creation and categorization, user messaging, notification history, and an administrative dashboard for managing users, alerts, and system records. The system utilizes a centralized database for storing tourist information, emergency notifications, and communication records. Furthermore, the developed system will be evaluated using the Functional Suitability characteristic of the ISO/IEC 25010 Software Quality Model to determine whether it performs its intended functions according to the specified requirements.",
        "The system enables authorized personnel to create, manage, broadcast, and resolve emergency alerts through Short Message Service (SMS), electronic mail (Email), and in-app (web channel) notifications. It includes functionalities such as tourist self-registration, secure user authentication, profile and notification preferences, the public Crisis Hub and resolved-alert archive, the Command Center for administrators, incident reporting with status tracking, route and travel advisories with maps, area hazard warnings, tourism guide dashboards and tour-group assignment, notification and activity history, CSV export of users and incidents, and administrative settings for broadcast audiences and channel testing. Active crisis alerts are visible on the Crisis Hub without login; SMS, email, and in-app delivery target registered users according to profile preferences and office broadcast rules. The system utilizes a centralized Supabase (PostgreSQL) database. Furthermore, the developed system will be evaluated using the Functional Suitability characteristic of the ISO/IEC 25010 Software Quality Model to determine whether it performs its intended functions according to the specified requirements.",
    ),
    (
        "The proposed system does not include disaster prediction, hazard monitoring, environmental sensing, geographic tracking, artificial intelligence-based decision support, or automatic emergency response dispatch.",
        "The proposed system does not include disaster prediction, automated environmental sensing, continuous real-time GPS tracking of tourists, artificial intelligence-based decision support, or automatic emergency response dispatch. Map features show office-published alert, route, and hazard locations but do not track devices in real time.",
    ),
    (
        "Without internet connectivity, users cannot register, log in, receive web notifications, send messages, or access emergency announcements through the system. Likewise, administrators cannot create emergency alerts, manage user accounts, monitor system activities, or respond to user messages while offline.",
        "Without internet connectivity, users cannot register, log in, receive in-app notifications, submit incident reports, or access the full web application. Likewise, administrators cannot create or broadcast emergency alerts, manage user accounts, or update incident report status while offline.",
    ),
    (
        "The messaging feature also depends on internet connectivity. Messages sent by registered users may not be delivered immediately if the user or the administrator experiences internet connection problems, power interruptions, server downtime, or hosting service interruptions. Messages will only be transmitted and received once connectivity and system services have been restored.",
        "Incident report submission and in-app notification delivery also depend on internet connectivity. Reports and status updates may not be processed immediately during connection problems, power interruptions, server downtime, or hosting service interruptions. Processing resumes once connectivity and system services are restored.",
    ),
    (
        "Furthermore, the system can only deliver notifications to registered users who provide valid mobile phone numbers and email addresses. Browser notifications also require users to enable notification permissions on their devices. The findings of this study are limited to the selected respondents and the testing environment used during the evaluation and may not necessarily represent the system's performance in all deployment environments.",
        "SMS and email notifications are sent only to registered users with valid contact information; SMS delivery additionally depends on user SMS preferences, the iProg SMS gateway, and mobile network coverage. In-app notifications require users to be signed in to the web application. The findings of this study are limited to the selected respondents and the testing environment used during the evaluation and may not necessarily represent the system's performance in all deployment environments.",
    ),
    (
        "Administrator – Refers to the authorized personnel of the Daet Tourism Office responsible for managing user accounts, creating and disseminating emergency alerts, responding to user messages, and maintaining the overall operation of the system.",
        "Administrator – Refers to the authorized personnel of the Daet Municipal Tourism Office responsible for managing user accounts, issuing and broadcasting emergency alerts through the Command Center, reviewing and updating incident reports, and maintaining the overall operation of the system.",
    ),
    (
        "Alert Category – Refers to the classification of emergency notifications according to the type of incident, such as typhoons, flooding, earthquakes, health emergencies, road accidents, or other public safety concerns.",
        "Alert Type and Severity – Refers to the classification of official crisis alerts (for example, typhoon, flooding, health advisory) together with severity levels (such as Critical, High, Medium, Low) assigned when an administrator publishes an alert.",
    ),
    (
        "Centralized Database – Refers to the repository used by the system to store and manage tourist information, emergency alerts, notification records, and user messages in a single database.",
        "Centralized Database – Refers to the Supabase (PostgreSQL) repository used to store and manage user profiles, emergency alerts, incident reports, notification and delivery records, route advisories, area hazards, and related tourism data in a single database.",
    ),
    (
        "Daet Tourism Crisis Communication and Emergency Alert System – Refers to the web-based information system developed in this study that enables the Daet Tourism Office to disseminate emergency notifications, manage tourist information, and facilitate communication between administrators and registered tourists.",
        "Daet Tourism Crisis Communication and Emergency Alert System (DAET TOURISM) – Refers to the web-based information system developed in this study that enables the Daet Municipal Tourism Office to publish crisis alerts, manage tourist and guide accounts, broadcast multi-channel notifications, and process incident reports from registered tourists.",
    ),
    (
        "Emergency Alert – Refers to an official notification created by authorized administrators to inform registered users about emergencies, hazards, or public safety incidents requiring immediate attention.",
        "Emergency Alert – Refers to an official crisis notification created by authorized administrators that appears on the public Crisis Hub when marked public and active, and may trigger SMS, email, and in-app notifications to registered users when broadcast.",
    ),
    (
        "User Messaging – Refers to the communication feature that allows registered users to send inquiries, reports, or concerns to the administrator through the web-based platform.",
        "Incident Report – Refers to the feature that allows registered tourists to submit hazard or emergency reports with category, severity, and location details, and to track tourism office response status through the system.",
    ),
    (
        "Web Notification – Refers to a browser-based notification generated by the system to provide real-time emergency alerts and announcements to users while they are logged into the web application.",
        "In-App Notification – Refers to a notification record delivered inside the signed-in DAET TOURISM application (Notifications inbox and on-screen alerts), distinct from SMS and email channels.",
    ),
    (
        "facilitating communication between tourism personnel and registered tourists.",
        "supporting incident reporting, guide monitoring of assigned tourists, and multi-channel alert delivery.",
    ),
    (
        "and user messaging within a single web-based platform specifically designed for the Daet Tourism Office and registered tourists.",
        "and incident reporting within a single web-based platform specifically designed for the Daet Municipal Tourism Office, registered tourists, and tourism guides.",
    ),
    (
        "and a user messaging feature to facilitate efficient communication between tourism personnel and registered tourists.",
        "and incident reporting with status tracking to facilitate structured communication between tourism personnel and registered tourists.",
    ),
    (
        "User Messaging",
        "Incident Reporting",
    ),
    (
        "It validates the use of the Next.js and Supabase tech stack for real-time synchronization.",
        "It validates the use of the Next.js, React, Tailwind CSS, Zustand, Leaflet, and Supabase tech stack for real-time synchronization and map-based advisories.",
    ),
    (
        "The conceptual design presents the overall structure of the proposed Daet Tourism Crisis Communication and Emergency Alert System. It illustrates the interaction between the Administrator, Registered Tourists, the web-based application, the centralized database, and the notification services. The conceptual design provides a high-level overview of how the system processes information and disseminates emergency alerts to registered users.",
        "The conceptual design presents the overall structure of the proposed Daet Tourism Crisis Communication and Emergency Alert System. It illustrates the interaction among Administrators, Registered Tourists, Tourism Guides, the web-based application (DAET TOURISM), the centralized Supabase database, and external notification services (SMS gateway and email). The conceptual design provides a high-level overview of how the system publishes public crisis information, broadcasts alerts to registered users, and processes incident reports.",
    ),
    (
        "Integrate SMS notifications, email notifications, web notifications, and user messaging.",
        "Integrate SMS (iProg), email (SMTP/Resend), in-app notifications, incident reporting, and Command Center broadcast workflows.",
    ),
    (
        "SMS integration, email integration, web notification integration, messaging module",
        "SMS integration, email integration, in-app notification integration, incident reporting module, route and hazard advisories",
    ),
    (
        "Develop the administrator dashboard, user management module, and alert category management.",
        "Develop the administrator dashboard, Command Center, user and guide management, and alert type/severity publishing.",
    ),
    (
        "Administrator dashboard, user management, profile management, alert category management",
        "Administrator dashboard, Command Center, user and guide management, profile and notification preferences",
    ),
    (
        "Alert Category Management",
        "Alert Type and Severity Classification",
    ),
    (
        "Web Notification System",
        "In-App Notification System",
    ),
    (
        "To analyze the existing emergency communication process of the Daet Tourism Office and identify the functional and non-functional requirements of the proposed system.",
        "To analyze the existing emergency communication process of the Daet Municipal Tourism Office and identify the functional and non-functional requirements of the proposed system.",
    ),
    (
        "organizing alert categories, monitoring notification history",
        "classifying alerts by type and severity, monitoring notification and audit history",
    ),
    (
        "create and categorize emergency alerts, disseminate notifications through SMS, email, and web notifications",
        "create and publish emergency alerts with type and severity, disseminate notifications through SMS, email, and in-app notifications",
    ),
    (
        "disseminate emergency alerts through SMS, email, and web notifications while facilitating communication between tourism personnel and registered tourists.",
        "disseminate emergency alerts through SMS, email, and in-app notifications while supporting incident reporting and guide monitoring for registered users.",
    ),
    (
        "Daet Tourism Office",
        "Daet Municipal Tourism Office",
    ),
]

NEW_DEFINITIONS = [
    "Tourism Guide – Refers to a user assigned the guide role by the Daet Municipal Tourism Office who can monitor crisis alerts, assigned tour groups, and tourists linked to their groups during an incident.",
    "Command Center – Refers to the administrative module where authorized personnel create, broadcast, update, and resolve official crisis alerts and review related operational activity.",
    "Crisis Hub – Refers to the public-facing section of DAET TOURISM where active and resolved official crisis alerts, maps, and safety instructions are displayed.",
]


def iter_all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def replace_in_paragraph(paragraph, old, new):
    if old not in paragraph.text:
        return False
    if "\n" in new and old in paragraph.text and paragraph.text.strip() == old:
        paragraph.text = new
        return True
    inline = paragraph.runs
    if not inline:
        paragraph.text = paragraph.text.replace(old, new)
        return True
    full = "".join(r.text for r in inline)
    if old not in full:
        return False
    full = full.replace(old, new.replace("\n", " "))
    for r in inline:
        r.text = ""
    inline[0].text = full
    return True


def insert_paragraph_after(paragraph, text):
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.text = text
    return new_para


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")
    if not BACKUP.exists():
        import shutil

        shutil.copy2(SRC, BACKUP)

    doc = Document(str(SRC))
    counts = {r[0][:40]: 0 for r in REPLACEMENTS}

    for old, new in REPLACEMENTS:
        for p in iter_all_paragraphs(doc):
            if replace_in_paragraph(p, old, new):
                counts[old[:40]] += 1

    # Fix accidental double "Municipal Municipal"
    for p in iter_all_paragraphs(doc):
        if "Municipal Municipal Tourism Office" in p.text:
            replace_in_paragraph(p, "Municipal Municipal Tourism Office", "Municipal Tourism Office")

    # Insert new definition terms before CHAPTER 2
    inserted = False
    for p in doc.paragraphs:
        if p.text.strip().startswith("In-App Notification –") and not inserted:
            anchor = p
            for line in NEW_DEFINITIONS:
                anchor = insert_paragraph_after(anchor, line)
            inserted = True
            break

    doc.save(str(OUT))
    applied = sum(1 for v in counts.values() if v > 0)
    print(f"Saved: {OUT}")
    print(f"Backup: {BACKUP}")
    print(f"Replacement patterns applied in at least one place: {applied}/{len(REPLACEMENTS)}")
    for k, v in counts.items():
        if v:
            print(f"  {v}x {k}...")


if __name__ == "__main__":
    main()
